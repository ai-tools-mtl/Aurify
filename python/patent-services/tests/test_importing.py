"""import round-trips: build a docx, import into a project, read the files back."""

from __future__ import annotations

from pathlib import Path

import pytest
import yaml
from docx import Document

from patent_services.importing import import_patent_document, split_document


def add_heading_like(document, text, *, bold=False):
    """Add one paragraph whose whole text is (optionally) bold — the shape the
    agency template's section headings take."""
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    return paragraph


def build_disclosure_docx(path: Path) -> Path:
    """An agency-template-shaped disclosure: numbered section headings, a
    bullet body, an unmapped 具体实施方式 section, and a claims section."""
    document = Document()
    add_heading_like(document, "一、名称：")
    document.add_paragraph("一种多智能体冲突消解方法")
    add_heading_like(document, "二、所属技术领域：")
    document.add_paragraph("本发明属于计算机软件技术领域。")
    add_heading_like(document, "三、背景技术：")
    document.add_paragraph("● 现有方案一：串行执行。")
    document.add_paragraph("综上，存在覆盖丢失问题。")
    add_heading_like(document, "四、现有技术的缺点是什么？针对这些缺点，说明本发明要解决的技术问题：")
    document.add_paragraph("1. 防止覆盖丢失。")
    add_heading_like(document, "五、发明内容（应该结合图形详细阐述该技术方案）：")
    document.add_paragraph("该方法增设四个协作组件。")
    add_heading_like(document, "六、有益效果")
    document.add_paragraph("消除静默覆盖丢失。")
    add_heading_like(document, "七、本发明的关键点和欲保护点是什么？")
    document.add_paragraph("关键点在意图声明与分层仲裁。")
    add_heading_like(document, "八、附图")
    document.add_paragraph("图1 为整体流程图。")
    add_heading_like(document, "九、具体实施方式")
    document.add_paragraph("以下结合附图说明实施例。")
    add_heading_like(document, "权利要求书")
    document.add_paragraph("1. 一种多智能体冲突消解方法，其特征在于……")
    document.save(path)
    return path


def test_split_document_maps_the_template_sections(tmp_path):
    sections, stats = build_and_split(tmp_path)
    kinds = {section.title: section.kind for section in sections}
    assert kinds["发明名称"] == "chapter"
    assert kinds["技术问题"] == "chapter"
    assert kinds["具体实施方式"] == "extra"
    assert kinds["权利要求书"] == "skip"
    assert kinds["关键点与保护范围"] == "chapter"
    solution = next(section for section in sections if section.title == "发明内容（技术方案）")
    assert "该方法增设四个协作组件。" in solution.lines


def build_and_split(tmp_path):
    source = build_disclosure_docx(tmp_path / "doc.docx")
    sections, stats = split_document(str(source))
    assert stats["tables"] == 0
    return sections, stats


def test_import_writes_chapters_brief_and_manifest(tmp_path):
    source = build_disclosure_docx(tmp_path / "existing.docx")
    project = tmp_path / "project"
    project.mkdir()
    report = import_patent_document(str(project), str(source))
    # The eight chapters all land — mapped with content, the rest placeholder.
    for filename in ("01-name.md", "02-field.md", "03-background.md", "04-problem.md",
                     "05-solution.md", "06-effect.md", "07-key-points.md", "08-drawings.md"):
        assert (project / "chapters" / filename).is_file(), filename
    name = (project / "chapters" / "01-name.md").read_text(encoding="utf-8")
    assert "一种多智能体冲突消解方法" in name
    extra = (project / "chapters" / "09-具体实施方式.md").read_text(encoding="utf-8")
    assert "以下结合附图说明实施例。" in extra
    # The claims section is named in the report, never written into chapters.
    assert "权利要求" in report and "未导入" in report
    claim_text = "\n".join(path.read_text(encoding="utf-8") for path in (project / "chapters").glob("*.md"))
    assert "其特征在于" not in claim_text
    # brief.md carries the five core sections seeded from the mapped bodies.
    brief = (project / "brief.md").read_text(encoding="utf-8")
    for heading in ("技术领域", "背景技术", "技术问题", "发明内容", "有益效果"):
        assert f"## {heading}" in brief
    assert "本发明属于计算机软件技术领域。" in brief
    # A minimal manifest appears so the loop's init gate passes.
    manifest = (project / "patent.yml").read_text(encoding="utf-8")
    assert "formatVersion: 1" in manifest and "status: drafting" in manifest
    assert "一种多智能体冲突消解方法" in manifest


def test_styled_subheadings_stay_inside_their_section(tmp_path):
    """Only a level-1 heading style splits chapters, and the level check runs
    BEFORE the keyword and extra-name matches: a ``Heading 2`` 组件设计 used
    to be shredded into ``09-``/``10-`` extra chapters; a ``Heading 2``
    技术方案细节 prefix-matched the solution keyword and lost its identity;
    a ``Heading 3`` 实施例一 prefix-matched the 实施例 extra name. All of
    them stay in the solution chapter as ``## ``/``### `` subheadings now."""
    document = Document()
    add_heading_like(document, "五、发明内容（应该结合图形详细阐述该技术方案）：")
    document.add_paragraph("总体架构分为两层。")
    document.add_heading("组件设计", level=2)
    document.add_paragraph("组件负责意图声明与冲突预检。")
    document.add_heading("接口设计", level=2)
    document.add_paragraph("接口走消息总线。")
    document.add_heading("技术方案细节", level=2)
    document.add_paragraph("细节展开如下。")
    document.add_heading("实施例一", level=3)
    document.add_paragraph("第一组实施例参数。")
    add_heading_like(document, "六、有益效果")
    document.add_paragraph("检索时延下降三成。")
    source = tmp_path / "styled.docx"
    document.save(str(source))
    project = tmp_path / "project"
    project.mkdir()
    import_patent_document(str(project), str(source))
    solution = (project / "chapters" / "05-solution.md").read_text(encoding="utf-8")
    assert "## 组件设计" in solution and "组件负责意图声明与冲突预检。" in solution
    assert "## 接口设计" in solution and "接口走消息总线。" in solution
    assert "## 技术方案细节" in solution and "细节展开如下。" in solution
    assert "### 实施例一" in solution and "第一组实施例参数。" in solution
    extras = sorted(path.name for path in (project / "chapters").glob("*.md")
                    if path.name[:2] >= "09")
    assert extras == []


def test_manifest_name_survives_yaml_metacharacters(tmp_path):
    """The manifest name is serialized, not interpolated: a name carrying
    ``: `` written raw produced a patent.yml every later read died on with a
    YAML ScannerError."""
    document = Document()
    add_heading_like(document, "一、名称：")
    document.add_paragraph("一种基于意图声明的校验方法")
    add_heading_like(document, "二、所属技术领域：")
    document.add_paragraph("数据处理。")
    source = tmp_path / "colon.docx"
    document.save(str(source))
    project = tmp_path / "project"
    project.mkdir()
    import_patent_document(str(project), str(source), name="方法: 系统")
    manifest = yaml.safe_load((project / "patent.yml").read_text(encoding="utf-8"))
    assert manifest["name"] == "方法: 系统"
    assert manifest["formatVersion"] == 1 and manifest["status"] == "drafting"


def test_import_bullets_and_subheadings_fold_into_markdown(tmp_path):
    document = Document()
    add_heading_like(document, "五、发明内容（应该结合图形详细阐述该技术方案）：")
    document.add_paragraph("● 意图声明组件：提交读写集。")
    sub = document.add_paragraph()
    sub.add_run("3.1 意图声明").bold = True
    document.add_paragraph("组件细节。")
    source = tmp_path / "doc.docx"
    document.save(source)
    project = tmp_path / "project"
    project.mkdir()
    import_patent_document(str(project), str(source))
    text = (project / "chapters" / "05-solution.md").read_text(encoding="utf-8")
    assert "- 意图声明组件：提交读写集。" in text
    assert "## 3.1 意图声明" in text


def test_import_refuses_nonempty_targets_without_overwrite(tmp_path):
    source = build_disclosure_docx(tmp_path / "existing.docx")
    project = tmp_path / "project"
    chapters = project / "chapters"
    chapters.mkdir(parents=True)
    (chapters / "03-background.md").write_text("# 背景技术\n\n已有内容，不许盲目覆盖。\n", encoding="utf-8")
    with pytest.raises(ValueError, match="overwrite"):
        import_patent_document(str(project), str(source))
    assert "已有内容" in (chapters / "03-background.md").read_text(encoding="utf-8")
    # overwrite=true goes through and replaces exactly the imported content.
    import_patent_document(str(project), str(source), overwrite=True)
    assert "综上，存在覆盖丢失问题。" in (chapters / "03-background.md").read_text(encoding="utf-8")


def test_import_keeps_existing_brief_and_manifest(tmp_path):
    source = build_disclosure_docx(tmp_path / "existing.docx")
    project = tmp_path / "project"
    project.mkdir()
    (project / "patent.yml").write_text("formatVersion: 1\nname: 既有项目\nstatus: review\n", encoding="utf-8")
    (project / "brief.md").write_text("# 技术交底摘要\n\n访谈所得，不能被导入覆盖。\n", encoding="utf-8")
    report = import_patent_document(str(project), str(source))
    assert "status: review" in (project / "patent.yml").read_text(encoding="utf-8")
    assert "访谈所得" in (project / "brief.md").read_text(encoding="utf-8")
    assert "brief.md 已按导入内容生成五维骨架" not in report
    assert "已建档" not in report


def test_import_fails_loud_on_unrecognizable_documents(tmp_path):
    project = tmp_path / "project"
    project.mkdir()
    document = Document()
    document.add_paragraph("这是一份与专利无关的会议纪要，没有任何章节标题。")
    source = tmp_path / "minutes.docx"
    document.save(source)
    with pytest.raises(ValueError, match="没有匹配到任何章节标题"):
        import_patent_document(str(project), str(source))

    non_docx = tmp_path / "scan.pdf"
    non_docx.write_bytes(b"%PDF-fake")
    with pytest.raises(ValueError, match="docx"):
        import_patent_document(str(project), str(non_docx))


def test_import_scaffolds_a_fresh_project_directory(tmp_path):
    """The import IS a project's starting point: a not-yet-created directory
    comes out as a full project (chapters + brief + manifest)."""
    source = build_disclosure_docx(tmp_path / "existing.docx")
    project = tmp_path / "fresh"
    report = import_patent_document(str(project), str(source))
    assert (project / "patent.yml").is_file()
    assert (project / "brief.md").is_file()
    assert (project / "chapters" / "05-solution.md").is_file()
    assert "已建档" in report


def test_import_names_the_prior_art_ledger_debt(tmp_path):
    document = Document()
    add_heading_like(document, "三、背景技术：")
    document.add_paragraph("现有方案见 CN109871542A。")
    source = tmp_path / "doc.docx"
    document.save(source)
    project = tmp_path / "project"
    project.mkdir()
    report = import_patent_document(str(project), str(source))
    assert "公开号" in report and "prior-art.md" in report
